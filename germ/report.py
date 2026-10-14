#!/usr/bin/env python3
"""Stage 4 - Excel germination report with standard seed-science metrics + curve."""
import json, argparse, numpy as np, pandas as pd
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows

FONT = "Arial"

def t50(times_h, counts, final):
    """Interpolated time to reach 50% of final germination count."""
    if final == 0:
        return None
    half = final / 2.0
    for i in range(1, len(counts)):
        if counts[i] >= half:
            c0, c1 = counts[i-1], counts[i]
            t0, t1 = times_h[i-1], times_h[i]
            if c1 == c0:
                return round(t1, 1)
            return round(t0 + (half - c0)*(t1 - t0)/(c1 - c0), 1)
    return None

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--map", default="seed_map.json")
    ap.add_argument("--per-seed", default="per_seed4.csv")
    ap.add_argument("--per-frame", default="per_frame4.csv")
    ap.add_argument("--out", default="germination_report.xlsx")
    args = ap.parse_args()

    cfg = json.loads(Path(args.map).read_text())
    dishes = [d["dish_id"] for d in cfg["dishes"]]
    n_per = {d["dish_id"]: d["n_seeds"] for d in cfg["dishes"]}
    ps = pd.read_csv(args.per_seed)
    pf = pd.read_csv(args.per_frame)

    # ---- per-dish summary metrics ----
    rows = []
    for d in dishes:
        sub = ps[ps.dish_id == d]
        n = n_per[d]
        g = int(sub.germinated.sum())
        gt = pd.to_numeric(sub[sub.germinated == 1]["germ_elapsed_h"], errors="coerce").dropna()
        mgt = float(gt.mean()) if len(gt) else np.nan
        counts = pf[f"{d}_germinated"].to_numpy()
        T50 = t50(pf.elapsed_h.to_numpy(), counts, g)
        rows.append(dict(
            Dish=d, Seeds_sown=n, Germinated=g,
            Germination_pct=round(100*g/n, 1),
            MGT_h=round(mgt, 1) if not np.isnan(mgt) else "",
            Mean_rate_per_h=round(1/mgt, 4) if (mgt and not np.isnan(mgt)) else "",
            T50_h=T50 if T50 is not None else "",
            First_germ_h=round(float(gt.min()), 1) if len(gt) else "",
            Last_germ_h=round(float(gt.max()), 1) if len(gt) else "",
        ))
    # overall row
    gall = int(ps.germinated.sum()); nall = int(sum(n_per.values()))
    gtall = pd.to_numeric(ps[ps.germinated == 1]["germ_elapsed_h"], errors="coerce").dropna()
    mgtall = float(gtall.mean()) if len(gtall) else np.nan
    T50all = t50(pf.elapsed_h.to_numpy(), pf.total_germinated.to_numpy(), gall)
    rows.append(dict(Dish="ALL", Seeds_sown=nall, Germinated=gall,
        Germination_pct=round(100*gall/nall, 1),
        MGT_h=round(mgtall, 1), Mean_rate_per_h=round(1/mgtall, 4),
        T50_h=T50all if T50all is not None else "",
        First_germ_h=round(float(gtall.min()), 1), Last_germ_h=round(float(gtall.max()), 1)))
    summ = pd.DataFrame(rows)

    # ---- build workbook ----
    wb = openpyxl.Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(name=FONT, bold=True, color="FFFFFF")
    title_font = Font(name=FONT, bold=True, size=14)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(ws, row, ncol):
        for c in range(1, ncol+1):
            cell = ws.cell(row=row, column=c)
            cell.fill = hdr_fill; cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    # Sheet 1: Summary
    ws = wb.active; ws.title = "Summary"
    ws["A1"] = "GermiCam — Seed Germination Report"; ws["A1"].font = title_font
    ws["A2"] = f"{cfg['n_dishes']} dishes · {cfg['n_seeds_total']} seeds · " \
               f"{len(pf)} timepoints · ~30 min cadence · germination = radicle/shoot emergence"
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    r0 = 4
    cols = list(summ.columns)
    for j, cname in enumerate(cols, 1):
        ws.cell(row=r0, column=j, value=cname.replace("_", " "))
    style_header(ws, r0, len(cols))
    for i, (_, rr) in enumerate(summ.iterrows(), 1):
        for j, cname in enumerate(cols, 1):
            cell = ws.cell(row=r0+i, column=j, value=rr[cname])
            cell.font = Font(name=FONT, bold=(rr["Dish"] == "ALL"))
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
    # column widths
    for j, cname in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = max(11, len(cname)+2)
    # notes
    nr = r0 + len(summ) + 3
    ws.cell(row=nr, column=1, value="Metric definitions").font = Font(name=FONT, bold=True)
    notes = [
        "Germination % (GP) = 100 × germinated / seeds sown.",
        "MGT (Mean Germination Time, h) = mean emergence time of germinated seeds = Σ(ni·ti)/Σni.",
        "Mean rate (1/h) = 1 / MGT.",
        "T50 (h) = time to reach 50% of the final germinated count (linear interpolation between frames).",
        "Detection is annotation-free (classical CV): fixed dish ROIs, orange-kernel seed map,",
        "sprout = green shoot OR new pale tissue connected to the kernel, latched over consecutive frames.",
        "Time axis derived from frame sequence at ~30 min cadence (no EXIF); calibrate if exact times needed.",
    ]
    for k, t in enumerate(notes):
        ws.cell(row=nr+1+k, column=1, value="• " + t).font = Font(name=FONT, size=9)

    # Sheet 2: Per_Seed
    ws2 = wb.create_sheet("Per_Seed")
    pcols = ["seed_id", "dish_id", "germinated", "germ_seq", "germ_elapsed_h",
             "final_len_r", "max_len_r"]
    for j, c in enumerate(pcols, 1):
        ws2.cell(row=1, column=j, value=c.replace("_", " "))
    style_header(ws2, 1, len(pcols))
    for i, (_, rr) in enumerate(ps[pcols].iterrows(), 2):
        for j, c in enumerate(pcols, 1):
            v = rr[c]
            ws2.cell(row=i, column=j, value=("" if pd.isna(v) else v))
    for j, c in enumerate(pcols, 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(j)].width = max(11, len(c)+2)
    ws2.freeze_panes = "A2"

    # Sheet 3: Per_Frame (+ chart source)
    ws3 = wb.create_sheet("Per_Frame")
    fcols = ["seq", "elapsed_h"] + [f"{d}_pct" for d in dishes] + ["total_pct"]
    hdrmap = {f"{d}_pct": f"{d} %" for d in dishes}; hdrmap["total_pct"] = "Total %"
    hdrmap["seq"] = "seq"; hdrmap["elapsed_h"] = "elapsed (h)"
    for j, c in enumerate(fcols, 1):
        ws3.cell(row=1, column=j, value=hdrmap[c])
    style_header(ws3, 1, len(fcols))
    for i, (_, rr) in enumerate(pf[fcols].iterrows(), 2):
        for j, c in enumerate(fcols, 1):
            ws3.cell(row=i, column=j, value=rr[c])
    for j, c in enumerate(fcols, 1):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 11
    ws3.freeze_panes = "A2"

    # chart: cumulative germination % vs time
    chart = LineChart()
    chart.title = "Cumulative germination (%) over time"
    chart.y_axis.title = "Germinated (%)"; chart.x_axis.title = "Elapsed time (h)"
    chart.height = 10; chart.width = 20
    nrows = len(pf) + 1
    data = Reference(ws3, min_col=3, max_col=len(fcols), min_row=1, max_row=nrows)
    cats = Reference(ws3, min_col=2, min_row=2, max_row=nrows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{nr+len(notes)+3}")

    wb.save(args.out)
    print("wrote", args.out)
    print(summ.to_string(index=False))

if __name__ == "__main__":
    main()
